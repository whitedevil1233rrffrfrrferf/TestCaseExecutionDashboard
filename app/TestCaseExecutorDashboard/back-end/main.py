import os
from fastapi import FastAPI, HTTPException, Query, WebSocket, BackgroundTasks,WebSocketDisconnect
from fastapi.responses import FileResponse
from services.ws_manager import ws_manager 
from openpyxl import Workbook
from typing import Optional, List
import tempfile
import os
import mysql.connector

import uvicorn
import asyncio
from mysql.connector import Error
from dotenv import load_dotenv
from fastapi.middleware.cors import CORSMiddleware
import sys
import randomname
from collections import defaultdict
from schemas import TestRunResponse,TestRunDetailsResponse,FilterResponse,AllFiltersResponse,TestRunSummaryResponse,TestRunFullResponse,RunEvaluationSummaryResponse,EvaluationItemResponse,ConversationResponse,TestCaseResponse,FullConversationResponse, TimelineEvent,NewTestRun
load_dotenv()
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..")))
import datetime
from datetime import datetime
from lib.orm import DB
from lib.data import Target, Run, RunDetail, Conversation

from lib.orm.tables import TestRuns
from lib.interface_manager import InterfaceManagerClient  # Import the InterfaceManagerClient from the lib directory

# db_url = (
#             f"mysql+mysqlconnector://"
#             f"{os.getenv('DB_USER')}:"
#             f"{os.getenv('DB_PASSWORD')}@"
#             f"{os.getenv('DB_HOST')}:"
#             f"{os.getenv('DB_PORT')}/"
#             f"{os.getenv('DB_NAME')}"
#         )

db_file = "AIEvaluationData.db"

# Resolve project root (this file → importer → app → src → project_root)
project_root = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../../.."))


# Place DB inside project_root/data
db_folder = os.path.join(project_root, "data")
os.makedirs(db_folder, exist_ok=True)

# Full DB path
db_path = os.path.join(db_folder, db_file)

# SQLite requires a file URL
db_url = f"sqlite:///{db_path}"

db = DB(db_url=db_url, debug=False)

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000"],  # React
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# def run_execute_testcases(*args):
#     import asyncio
#     asyncio.run(execute_testcases(*args))

## WebSocket for real-time updates

@app.websocket("/ws/test-run")
async def websocket_endpoint(websocket: WebSocket):
    await ws_manager.connect(websocket)
    try:
        while True:
            # keep connection alive
            await websocket.receive_text()
    except WebSocketDisconnect:
        ws_manager.disconnect(websocket)

async def step(ws_payload, delay=0.1):
    await ws_manager.send_all(ws_payload)
    await asyncio.sleep(delay)

@app.get(
    "/get_all_test_runs",
    response_model=list[TestRunResponse]
)
def get_all_test_runs(
    domain: Optional[str] = Query(None),
    target: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
):
    try:
        
        db = DB(db_url=db_url, debug=False)
        runs = db.get_all_runs(domain=domain, target=target, status=status)
        print(domain, target, status)
        response = []

        for r in runs:
            domain_name = None

            target_id = r.kwargs.get("target_id") if hasattr(r, "kwargs") else None

            if target_id:
                target_obj = db.get_target_by_id(target_id)
                if target_obj:
                    domain_name = target_obj.target_domain   # ✅ FIX

            response.append(
                TestRunResponse(
                    run_id=r.run_id,
                    run_name=r.run_name,
                    target=r.target,
                    status=r.status,
                    start_ts=r.start_ts,
                    end_ts=r.end_ts,
                    domain=domain_name
                )
            )

        return response

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    
@app.get(
    "/test-runs/{run_name}",
    response_model=TestRunFullResponse
)
def get_test_run(
    run_name: str,
    metric: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    ):
    try:
        
        print(metric, status)
        # ---------- RUN SUMMARY ----------
        run = db.get_run_by_name(run_name)
        if not run:
            raise HTTPException(status_code=404, detail="Run not found")

        domain_name = None
        if getattr(run, "target_id", None):
            target = db.get_target_by_id(run.target_id)
            if target:
                domain_name = getattr(target, "target_domain", None)

        summary = TestRunSummaryResponse(
            run_id=run.run_id,
            run_name=run.run_name,
            target=run.target,
            domain=domain_name,
            status=run.status,
            start_ts=run.start_ts,
            end_ts=run.end_ts
        )

        # ---------- RUN DETAILS ----------
        details = db.get_all_run_details_by_run_name(run_name)
        
        details_response = []
        if metric:
            details = [d for d in details if d.metric_name == metric]

        if status:
            details = [d for d in details if d.status == status]
            
        for d in details:
            score = None

            if d.conversation_id:
                conv = db.get_conversation_by_id(d.conversation_id)
                if conv and conv.evaluation_score is not None:
                    score = float(conv.evaluation_score)

            details_response.append(
                TestRunDetailsResponse(
                    run_name=d.run_name,
                    testcase_name=d.testcase_name,
                    metric_name=d.metric_name,
                    plan_name=d.plan_name,
                    conversation_id=d.conversation_id,
                    status=d.status,
                    detail_id=d.detail_id,
                    score=score
                )
            )

        # ---------- FINAL RESPONSE ----------
        return TestRunFullResponse(
            summary=summary,
            details=details_response
        )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/get_all_filters", response_model=AllFiltersResponse)
def get_all_filters():
    try:
        

        # Use the @property methods to get all data
        return AllFiltersResponse(
            domains=[FilterResponse(filter_name=d.name) for d in db.domains],
            languages=[FilterResponse(filter_name=l.name) for l in db.languages],
            targets=[FilterResponse(filter_name=t.target_name) for t in db.targets],
            statuses=[
                FilterResponse(filter_name="COMPLETED"),
                FilterResponse(filter_name="RUNNING"),
            ],
            plans=[FilterResponse(filter_name=p.plan_name) for p in db.plans],
            metrics=[FilterResponse(filter_name=m.metric_name) for m in db.metrics]
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    
@app.get("/test-runs/{run_name}/summary", response_model=TestRunSummaryResponse)
def get_test_run_summary(run_name: str):
    try:
        

        run = db.get_run_by_name(run_name)
        if not run:
            raise HTTPException(status_code=404, detail="Run not found")

        # If you store target_id on run, use it to fetch target -> domain
        domain_name = None
        if getattr(run, "target_id", None):
            target = db.get_target_by_id(run.target_id)
            if target:
                domain_name = getattr(target, "target_domain", None)

        return TestRunSummaryResponse(
            run_id=run.run_id,
            run_name=run.run_name,
            target=run.target,
            domain=domain_name,
            status=run.status,
            start_ts=run.start_ts,
            end_ts=run.end_ts
        )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))    
    
@app.get(
    "/test-runs/{run_name}/evaluation-summary",
    response_model=RunEvaluationSummaryResponse
)
def get_run_evaluation_summary(run_name: str):
    try:
       
        run= db.get_run_by_name(run_name)
        if not run:
            raise HTTPException(status_code=404, detail="Run not found")
        domain_name = None
        if getattr(run, "target_id", None):
            target = db.get_target_by_id(run.target_id)
            if target:
                domain_name = getattr(target, "target_domain", None)
        run_summary = TestRunSummaryResponse(
            run_id=run.run_id,
            run_name=run.run_name,
            target=run.target,
            domain=domain_name,
            status=run.status,
            start_ts=run.start_ts,
            end_ts=run.end_ts
        )        
        details = db.get_all_run_details_by_run_name(run_name)

        evaluations=[]

        for d in details:
            conv = db.get_conversation_by_id(d.conversation_id)
            print(conv.evaluation_score)
            if not conv:
                continue
            evaluations.append(
                EvaluationItemResponse(
                    detail_id=d.detail_id,
                    testcase=conv.testcase,
                    agent_response=conv.agent_response,
                    evaluation_score=conv.evaluation_score,
                    evaluation_reason=conv.evaluation_reason,
                    evaluation_ts=conv.evaluation_ts
                )
            )

        return RunEvaluationSummaryResponse(
            run=run_summary,
            evaluations=evaluations
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
   

@app.get("/test-runs/{run_name}/evaluation-report")
def download_evaluation_report(run_name: str):
    try:
        

        # -------------------------------------------------
        # FETCH RUN
        # -------------------------------------------------
        run = db.get_run_by_name(run_name)
        if not run:
            raise HTTPException(status_code=404, detail="Run not found")

        details = db.get_all_run_details_by_run_name(run_name)
        plan_name = details[0].plan_name if details else None
        # Cache conversations
        conversation_cache = {}
        testcase_prompt_cache = {}
        for d in details:
            if d.conversation_id not in conversation_cache:
                conversation_cache[d.conversation_id] = db.get_conversation_by_id(
                    d.conversation_id
                )

        wb = Workbook()

        # =================================================
        # SHEET 1 : RUN SUMMARY (NO BS)
        # =================================================
        ws_summary = wb.active
        ws_summary.title = "Run_Summary"

        # Collect counts
        testcases = set()
        total_metrics = len(details)

        for conv in conversation_cache.values():
            if conv and getattr(conv, "testcase", None):
                testcases.add(conv.testcase)

        ws_summary.append(["Run Name", run.run_name])
        ws_summary.append(["Plan Name", plan_name])
        ws_summary.append(["Status", run.status])
        ws_summary.append(["Total Testcases", len(testcases)])
        ws_summary.append(["Total Metrics", total_metrics])

        # =================================================
        # SHEET 2 : EVALUATION DETAILS (TESTCASE + METRIC)
        # =================================================
        ws_details = wb.create_sheet(title="Evaluation_Details")

        ws_details.append([
            "Detail ID",
            "Testcase",
            "Metric",
            "Evaluation Score",
            "Evaluation Reason",
            "Evaluation Time",
            "Agent Response",
            "User Prompt",
            "System Prompt"
        ])

        for d in details:
            conv = conversation_cache.get(d.conversation_id)
            if not conv:
                continue

            metric_name = (
                getattr(d, "metric", None)
                or getattr(d, "metric_name", None)
                or getattr(conv, "metric", None)
                or getattr(conv, "metric_name", None)
            )
            testcase_name = getattr(conv, "testcase", None)
            if testcase_name not in testcase_prompt_cache:
                testcase = db.get_testcase_by_name(testcase_name)

                if not testcase or not getattr(testcase, "prompt", None):
                    testcase_prompt_cache[testcase_name] = {
                        "user_prompt": None,
                        "system_prompt": None,
                    }
                else:
                    testcase_prompt_cache[testcase_name] = {
                        "user_prompt": getattr(testcase.prompt, "user_prompt", None),
                        "system_prompt": getattr(testcase.prompt, "system_prompt", None),
                    }
            prompts = testcase_prompt_cache[testcase_name]        
            ws_details.append([
                d.detail_id,
                getattr(conv, "testcase", None),
                metric_name,
                getattr(conv, "evaluation_score", None),
                getattr(conv, "evaluation_reason", None),
                getattr(conv, "evaluation_ts", None),
                getattr(conv, "agent_response", None),
                prompts["user_prompt"],
                prompts["system_prompt"],
            ])

        # =================================================
        # SAVE FILE
        # =================================================
        tmp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        wb.save(tmp_file.name)
        tmp_file.close()

        return FileResponse(
            path=tmp_file.name,
            filename=f"{run_name}_evaluation_report.xlsx",
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    
@app.get(
    "/testcases/{testcase_name}",
    response_model=TestCaseResponse
)
def get_conversation(testcase_name: str):
    
    testcase = db.get_testcase_by_name(testcase_name)
    if not testcase:
        raise HTTPException(status_code=404, detail="Testcase not found")
    return TestCaseResponse(
        user_prompt=testcase.prompt.user_prompt,
        system_prompt=testcase.prompt.system_prompt
    )

@app.get(
    "/conversations/full/{conversation_id}",
    response_model=FullConversationResponse
)
def get_full_conversation(conversation_id: int):
    
    conversation = db.get_conversation_by_id(conversation_id)
    if not conversation:
        raise HTTPException(status_code=404, detail="Conversation not found")

    testcase_name = conversation.testcase
    testcase = db.get_testcase_by_name(testcase_name)
    if not testcase:
        user_prompt = None
        system_prompt = None
    else:
        user_prompt = getattr(testcase.prompt, "user_prompt", None)
        system_prompt = getattr(testcase.prompt, "system_prompt", None)

    return FullConversationResponse(
        user_prompt=user_prompt,
        system_prompt=system_prompt,
        agent_response=conversation.agent_response,
        testcase_name=testcase_name,
        conversation_id=conversation_id,
        target=conversation.target,
        score=conversation.evaluation_score,
        reason=conversation.evaluation_reason
    )

@app.get("/conversations/{conversation_id}/timeline")
def get_conversation_timeline_api(conversation_id: int):
    
    timeline = db.get_conversation_timeline(conversation_id)

    if not timeline:
        raise HTTPException(status_code=404, detail="Conversation not found")

    return timeline

@app.get(
    "/test-runs/{run_name}/timeline",
    response_model=list[TimelineEvent]
)
def get_test_run_timeline(run_name: str):
    
    timeline = db.get_run_timeline(run_name)
    if not timeline:
        raise HTTPException(status_code=404, detail="No timeline found")
    
    return timeline

@app.post("/start-run")
def start_run(data: NewTestRun, background_tasks: BackgroundTasks):
    if data.testPlanId:
        ### Initialising the form variables

        target = data.target
        test_plan_id = data.testPlanId
        test_case_id = data.testCaseId
        metric_id = data.metricId
        lang_names=None
        domain_name=None

        ## Create a random name for the run and generating run id

        run_name = randomname.generate('v/*','adj/*','n/*','ip/*')
        start_time = datetime.now().isoformat()
        run = Run(target=target, run_name=run_name, start_ts=start_time)
        run_id = db.add_or_update_testrun(run)
        print(f"Starting run: {run_name} with run id {run_id}")

        
        plan_name = db.get_testplan_name(plan_id=test_plan_id)
        
        if plan_name is None:
            print(f"No test plan found with ID {test_plan_id}.")
            return
        print(f"Starting run with Test Plan: {plan_name} (ID: {test_plan_id})")

        if metric_id:
            metric_name = db.get_metric_name(metric_id=metric_id)  
            is_metric_in_plan = db.is_metric_in_testplan(metric_name=metric_name, plan_name=plan_name)  
            if not is_metric_in_plan:
                return
            testcases = db.get_testcases_by_metric(
                metric_name=db.get_metric_name(metric_id),
                n=3,
                lang_name=None,
                domain_name=None
            ) 
            if not testcases:
                print("No Test cases Found")
                return
            #     ## Get the metric from the provided ID 
            total_testcases = len(testcases)
            metric = db.get_metric_by_id(metric_id=metric_id)
            if metric is None:
                print("No Metric Found")
                return
            run.status = "RUNNING"
            db.add_or_update_testrun(run=run)
            background_tasks.add_task(
                execute_testcases,
                run_name,
                run_id,
                plan_name,
                target,
                metric_id,
                testcases,
                run
            )
            # agent_name = target
            # application_name = target
            # application_url = "https://web.whatsapp.com"
            # application_type = "WHATSAPP_WEB"
            # print("started syncing")
            # client = InterfaceManagerClient(base_url="http://localhost:8000" ,application_type=application_type, agent_name=agent_name)
            # client.sync_config({
            #         "application_name": application_name,
            #         "application_type": application_type,
            #         "agent_name": agent_name,
            #         "application_url": application_url
            #     })
            # client.apply_server_config()
            # for testcase in testcases:
            #     rundetail = RunDetail(run_name=run_name, plan_name=plan_name, metric_name=testcase.metric, testcase_name=testcase.name)
            #     rundetail_id = db.add_or_update_testrun_detail(rundetail)
            #     run_status = db.get_status_by_run_detail_id(run_detail_id=rundetail_id)
            #     if run_status is not None and run_status == "COMPLETED":
            #         print(f"Run detail for testcase {testcase.name} (ID: {testcase.testcase_id}) is already completed. Skipping execution.")
            #         continue
            #     message_to_agent = testcase.prompt.user_prompt if testcase.prompt.user_prompt else ""
            #     if testcase.prompt.system_prompt:
            #         message_to_agent = testcase.prompt.system_prompt + " " + message_to_agent

            #     conv = Conversation(target=target, 
            #                         run_detail_id=rundetail_id, 
            #                         testcase=testcase.name)
            #     conv_id = db.add_or_update_conversation(conversation=conv)
            #     print(f"A new conversation is created with ID: {conv_id}")
            #     rundetail.status = "RUNNING"
            #     db.add_or_update_testrun_detail(rundetail)
            #     # print("completed")
                
            #     conv.prompt_ts = datetime.now().isoformat()
            #     db.add_or_update_conversation(conversation=conv)
            #     print("prompt time added")    
            #     response_from_agent = client.chat(chat_id = testcase.testcase_id, prompt_list=[message_to_agent])
            #     agent_response = response_from_agent.json().get("response", "")
            #     if len(agent_response) == 0 or agent_response[0]['response'] == "Chat not found":
            #         print(f"No response received from the agent for test case {testcase.testcase_id}.")
            #         rundetail.status = "FAILED"
            #         db.add_or_update_testrun_detail(rundetail)
            #         continue
            #     conv.response_ts = datetime.now().isoformat()
            #     print("Response time added") 
            #     conv.agent_response = agent_response[0]['response']
            #     db.add_or_update_conversation(conversation=conv)
            #     rundetail.status = "COMPLETED"
            #     db.add_or_update_testrun_detail(rundetail)
            #     print("completed Response time")
            
        # 🔹 Get testcases count (NO execution here)

        return {
            "status": "success",
            "runId": run_id,
            "runName": run_name,
            "testPlanId": test_plan_id,
            "metricId": metric_id,
            "target": target,
            "totalTestCases": total_testcases,
            
        }
        ## Create a random name for the run and generating run id

        # run_name = randomname.generate('v/*','adj/*','n/*','ip/*')
        # start_time = datetime.now().isoformat()
        # print(type(start_time))
        # run = Run(target = target, run_name=run_name, start_ts=start_time)
        # # print(run)
        # run_id = db.add_or_update_testrun(run=run)
        # print(f"Starting run: {run_name} with run id {run_id}")

        # ## Getting the plan name 

        # plan_name = db.get_testplan_name(plan_id=test_plan_id)

        # if plan_name is None:
        #     print(f"No test plan found with ID {test_plan_id}.")
        #     return
        # print(f"Starting run with Test Plan: {plan_name} (ID: {test_plan_id})")

        # ## checking the conditionality for metric id and test case id provided and updating the test run

        # if test_case_id:
        #     testcase = db.get_testcase_by_id(testcase_id=test_case_id)
        #     if testcase is None:
        #             print(f"No test case found with ID {test_case_id}.")
        #             return
        #     run.status = "RUNNING"
        #     db.add_or_update_testrun(run=run)
        #     rundetail = RunDetail(run_name=run_name, plan_name=plan_name, metric_name=testcase.metric, testcase_name=testcase.name)
        #     rundetail_id = db.add_or_update_testrun_detail(rundetail)
        #     run_status = db.get_status_by_run_detail_id(run_detail_id=rundetail_id)
        #     if run_status is not None and run_status == "COMPLETED":
        #         print("Already completed")

        #     else:
        #         print("Sucessfully done with TestCaseId")  

        # if metric_id:

        #     ## Get the metric name and see if it matches with metric id and test plan id

        #     metric_name = db.get_metric_name(metric_id=metric_id)  
        #     is_metric_in_plan = db.is_metric_in_testplan(metric_name=metric_name, plan_name=plan_name)  
        #     if not is_metric_in_plan:
        #         print("error")


        #     testcases = db.get_testcases_by_metric(metric_name=metric_name, n=0, lang_name=lang_names, domain_name=domain_name)    

        #     if not testcases:
        #         print("No Test cases Found")
        #         return

        #     ## Get the metric from the provided ID 

        #     metric = db.get_metric_by_id(metric_id=metric_id)
        #     if metric is None:
        #         print("No Metric Found")
        #         return
        #     run.status = "RUNNING"
        #     db.add_or_update_testrun(run=run)
        #     agent_name = target
        #     application_name = target
        #     application_url = "https://web.whatsapp.com"
        #     application_type = "WHATSAPP_WEB"
        #     print("started syncing")
        #     client = InterfaceManagerClient(base_url="http://localhost:8000" ,application_type=application_type, agent_name=agent_name)
        #     client.sync_config({
        #             "application_name": application_name,
        #             "application_type": application_type,
        #             "agent_name": agent_name,
        #             "application_url": application_url
        #         })
        #     client.apply_server_config()
        #     for testcase in testcases:
        #         rundetail = RunDetail(run_name=run_name, plan_name=plan_name, metric_name=testcase.metric, testcase_name=testcase.name)
        #         rundetail_id = db.add_or_update_testrun_detail(rundetail)
        #         run_status = db.get_status_by_run_detail_id(run_detail_id=rundetail_id)
        #         if run_status is not None and run_status == "COMPLETED":
        #             print(f"Run detail for testcase {testcase.name} (ID: {testcase.testcase_id}) is already completed. Skipping execution.")
        #             continue
        #         message_to_agent = testcase.prompt.user_prompt if testcase.prompt.user_prompt else ""
        #         if testcase.prompt.system_prompt:
        #             message_to_agent = testcase.prompt.system_prompt + " " + message_to_agent

        #         conv = Conversation(target=target, 
        #                             run_detail_id=rundetail_id, 
        #                             testcase=testcase.name)
        #         conv_id = db.add_or_update_conversation(conversation=conv)
        #         print(f"A new conversation is created with ID: {conv_id}")
        #         rundetail.status = "RUNNING"
        #         db.add_or_update_testrun_detail(rundetail)
        #         # print("completed")
                
        #         conv.prompt_ts = datetime.now().isoformat()
        #         db.add_or_update_conversation(conversation=conv)
        #         print("prompt time added")    
        #         response_from_agent = client.chat(chat_id = testcase.testcase_id, prompt_list=[message_to_agent])
        #         agent_response = response_from_agent.json().get("response", "")
        #         if len(agent_response) == 0 or agent_response[0]['response'] == "Chat not found":
        #             print(f"No response received from the agent for test case {testcase.testcase_id}.")
        #             rundetail.status = "FAILED"
        #             db.add_or_update_testrun_detail(rundetail)
        #             continue
        #         conv.response_ts = datetime.now().isoformat()
        #         print("Response time added") 
        #         conv.agent_response = agent_response[0]['response']
        #         db.add_or_update_conversation(conversation=conv)
        #         rundetail.status = "COMPLETED"
        #         db.add_or_update_testrun_detail(rundetail)
        #         print("completed Response time")
                
        return {"status": "success"}
    else:
        return("Test Plan ID is mandatory")



async def execute_testcases(
    run_name,
    run_id,
    plan_name,
    target,
    metric_id,
    testcases,
    run
):
    print(f"🚀 Background execution started for run {run_id}")

    agent_name = target
    application_name = target
    application_url = "https://web.whatsapp.com"
    application_type = "WHATSAPP_WEB"

    client = InterfaceManagerClient(
        base_url="http://localhost:8000",
        application_type=application_type,
        agent_name=agent_name
    )
    await ws_manager.send_all({
        "type": "RUN_STARTED",
        "runId": run_id,
        "total": len(testcases)
    })
    client.sync_config({
        "application_name": application_name,
        "application_type": application_type,
        "agent_name": agent_name,
        "application_url": application_url
    })
    client.apply_server_config()

    for index, testcase in enumerate(testcases, start=1):
        print(f"⚙️ Running testcase: {testcase.name}")

        rundetail = RunDetail(
            run_name=run_name,
            plan_name=plan_name,
            metric_name=testcase.metric,
            testcase_name=testcase.name
        )
        rundetail_id = db.add_or_update_testrun_detail(rundetail)
        run_status = db.get_status_by_run_detail_id(run_detail_id=rundetail_id)
        if run_status is not None and run_status == "COMPLETED":
            print(f"Run detail for testcase {testcase.name} (ID: {testcase.testcase_id}) is already completed. Skipping execution.")
            continue

        message_to_agent = testcase.prompt.user_prompt or ""
        if testcase.prompt.system_prompt:
            message_to_agent = testcase.prompt.system_prompt + " " + message_to_agent

        conv = Conversation(target=target, 
                            run_detail_id=rundetail_id, 
                            testcase=testcase.name)
        conv_id = db.add_or_update_conversation(conversation=conv)    
        print(f"A new conversation is created with ID: {conv_id}")

        # await ws_manager.send_all({
        #     "type": "STEP_UPDATE",
        #     "runId": run_id,
        #     "testcaseIndex": index,
        #     "step": 1,
        #     "status": "RUNNING"
        # })
        rundetail.status = "RUNNING"
        db.add_or_update_testrun_detail(rundetail)
        conv.prompt_ts = datetime.now().isoformat()
        db.add_or_update_conversation(conversation=conv)

        await step({
            "type": "STEP_UPDATE",
            "runId": run_id,
            "testcaseIndex": index,
            "step": 1,
            "status": "DONE"
        })
        # await ws_manager.send_all({
        #     "type": "STEP_UPDATE",
        #     "runId": run_id,
        #     "testcaseIndex": index,
        #     "step": 2,
        #     "status": "RUNNING"
        # })
        response_from_agent = client.chat(
            chat_id=testcase.testcase_id,
            prompt_list=[message_to_agent]
        )
        await step({
            "type": "STEP_UPDATE",
            "runId": run_id,
            "testcaseIndex": index,
            "step": 2,
            "status": "DONE"
        })
        # await ws_manager.send_all({
        #     "type": "STEP_UPDATE",
        #     "runId": run_id,
        #     "testcaseIndex": index,
        #     "step": 3,
        #     "status": "RUNNING"
        # })
        agent_response = response_from_agent.json().get("response", "")
        if len(agent_response) == 0 or agent_response[0]['response'] == "Chat not found":
            print(f"No response received from the agent for test case {testcase.testcase_id}.")
            rundetail.status = "FAILED"
            db.add_or_update_testrun_detail(rundetail)
            continue
        conv.response_ts = datetime.now().isoformat()
        conv.agent_response = agent_response[0]['response']
        db.add_or_update_conversation(conversation=conv)

        await step({
            "type": "STEP_UPDATE",
            "runId": run_id,
            "testcaseIndex": index,
            "step": 3,
            "status": "DONE"
        })
        # await ws_manager.send_all({
        #     "type": "STEP_UPDATE",
        #     "runId": run_id,
        #     "testcaseIndex": index,
        #     "step": 4,
        #     "status": "RUNNING"
        # })
        await step({
        "type": "STEP_UPDATE",
        "runId": run_id,
        "testcaseIndex": index,
        "step": 4,
        "status": "DONE"
        })
        await step({
        "type": "TESTCASE_FINISHED",
        "runId": run_id,
        "current": index
        })
    rundetail.status = "COMPLETED"
    db.add_or_update_testrun_detail(rundetail)
    run.end_ts = datetime.now().isoformat()
    run.status = "COMPLETED"
    db.add_or_update_testrun(run=run)
    
        # client.close()
    
   

    print(f"✅ Finished testcase: {testcase.name}")
    await ws_manager.send_all({
        "type": "RUN_FINISHED",
        "runId": run_id
    })
    print(f"🏁 Background execution finished for run {run_id}")




if __name__ == "__main__":
    uvicorn.run(
        "main:app",
        host="0.0.0.0",
        port=7000,
        reload=True
    )